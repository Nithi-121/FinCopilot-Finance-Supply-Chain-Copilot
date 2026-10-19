"""Build the complete FinCopilot project walkthrough and source appendix."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "FinCopilot_Data_Science_Project_Report.docx"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(40, 48, 58)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in [("Title", 29, "14324A"), ("Heading 1", 19, "14324A"), ("Heading 2", 13, "147D92"), ("Heading 3", 11, "275D72")]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("FinCopilot | Data Science Project Report")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(105, 115, 125)


def para(doc, text: str, style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def bullet(doc, text: str):
    return para(doc, text, "List Bullet")


def number(doc, text: str):
    return para(doc, text, "List Number")


def table(doc, headers: list[str], rows: list[list[str]]):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    for i, title in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = title
        shade(cell, "14324A")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()
    return t


def code_listing(doc, path: Path):
    doc.add_heading(path.relative_to(ROOT).as_posix(), level=2)
    source = path.read_text(encoding="utf-8").splitlines()
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(source):
        r = p.add_run(line + ("\n" if i < len(source) - 1 else ""))
        r.font.name = "Cascadia Mono"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Cascadia Mono")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Cascadia Mono")
        # The final test module narrowly exceeded one page at the default code size.
        r.font.size = Pt(7.0 if path.name == "test_core.py" else 7.3)
        r.font.color.rgb = RGBColor(45, 55, 65)


def main():
    doc = Document()
    setup(doc)
    title = doc.add_paragraph(style="Title")
    title.add_run("FinCopilot Data Science Project Report")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    sr = subtitle.add_run("A filing research assistant with hybrid retrieval, verified citations, and finance data tools")
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor(90, 104, 115)
    para(doc, "Project 3 | Finance and data science portfolio build | 17 October 2026")

    doc.add_heading("Executive summary", level=1)
    para(doc, "FinCopilot is a reproducible data science project for asking questions about annual reports from HPE, Microsoft, Apple, NVIDIA, and Accenture. It retrieves filing passages with a combination of local semantic embeddings and BM25 keyword search, shows source metadata, and uses a configurable Ollama local or cloud model to write a plain-language response. Financial values come from read-only DuckDB tools populated from SEC XBRL company facts; an optional FinPulse connection can replace this local database.")
    para(doc, "The project includes an SEC EDGAR fetcher, filing parser and chunker, Chroma vector index, three-mode retriever with rank diagnostics, citation checks, read-only DuckDB tools, exploratory forecasting and anomaly analysis, a conversational Ollama assistant, FastAPI endpoints, a Streamlit interface, and this standalone responsive dashboard. The dashboard supports chat history, company comparisons, source and tool inspection, analytics, and exports. SEC contact details stay in the ignored .env file; the local showcase server serves only the HTML page.")
    para(doc, "Project status at report generation: 15 public SEC filings (three each for five companies) have been downloaded, recorded with SHA-256 hashes, and parsed into 1,519 source-tagged chunks across the four planned filing sections. SEC company facts produced 644 historical and derived metric rows, 36 one-year forecast rows, and 34 candidate outlier rows in local DuckDB. The Chroma index is built; all fifteen offline tests and Ruff checks pass. The signed-in Ollama 0.40.5 service is configured for gemma4:cloud; plain-language finance definitions and API readiness were verified, and the read-only analytics endpoint returned the local dataset rows. A preliminary section-level seed comparison measured hit@5 of 0.825 for BM25 and 0.925 for dense and hybrid retrieval; these labels are not verified and are not a final benchmark. Citation validity, refusal scoring, and answer quality remain unmeasured. The earlier FinPulse database is not present, so its optional external connection is not tested.")

    doc.add_heading("Contents", level=1)
    for entry in [
        "1 Purpose and data science role fit", "2 User question and project outcome",
        "3 Architecture and information flow", "4 Build steps in order",
        "5 Module guide", "6 Dataset and evaluation design",
        "7 Data quality, reliability, and safety", "8 Run commands and expected artifacts",
        "9 Portfolio description guidance", "Appendix A Complete source code",
    ]:
        bullet(doc, entry)

    doc.add_heading("1. Purpose and data science role fit", level=1)
    para(doc, "This project demonstrates work that maps to a data scientist role: collecting semi-structured public data; cleaning and segmenting documents; defining an information retrieval problem; building an indexed dataset; designing an evaluation set; measuring retrieval quality; conducting ablation comparisons; connecting model outputs to structured data; and communicating failure cases and limits.")
    table(doc, ["Skill area", "Evidence in this project"], [
        ["Data acquisition", "15 SEC EDGAR 10-K filings and public XBRL company facts across five companies"],
        ["Data preparation", "HTML cleanup, section segmentation, 1,519 overlapping text chunks with source metadata"],
        ["Retrieval modeling", "Sentence-transformer embeddings, Chroma, BM25, reciprocal-rank fusion"],
        ["Evaluation", "40 answerable and 10 unanswerable seed prompts; BM25/dense/hybrid hit@k harness; human gold-label workflow"],
        ["Applied ML systems", "Ollama local/cloud LLM orchestration, grounded citations, DuckDB tool calls, API and user interface"],
        ["Time series", "36 annual baseline forecasts selected through walk-forward error comparison"],
        ["Anomaly detection", "34 candidate flags from unusual year-over-year changes across annual metrics"],
        ["Responsible analysis", "No invented benchmark scores; local secret handling; read-only SQL; explicit refusal and advisory limits"],
    ])
    para(doc, "The five-company universe includes Accenture (ACN) as requested. The company list can be edited through TICKERS in .env without changing the SEC CIK lookup code.")

    doc.add_heading("2. User question and project outcome", level=1)
    para(doc, "The use case is a question such as: ‘What risks did Accenture describe, and how did its operating margin change?’ The assistant should retrieve relevant Item 1A and Item 7 passages, cite each filing claim, and request numeric values from the connected financial database. If the supporting evidence or database value is unavailable, it should say so rather than guessing.")
    para(doc, "The generated SEC database provides reported history and a one-step annual baseline forecast. The previous FinPulse database is not in the workspace; it can replace the SEC-built DuckDB after its path and schema are configured.")

    doc.add_heading("3. Architecture and information flow", level=1)
    table(doc, ["Stage", "Input", "Processing", "Output"], [
        ["Collect", "Five configured tickers", "Resolve CIKs from SEC ticker map; download 10-K HTML and XBRL company facts", "15 local filings, SEC facts JSON, and SHA-256 filing manifest"],
        ["Prepare", "SEC HTML", "Remove non-content markup; detect filing sections; make 650-word chunks with 100-word overlap", "JSONL chunks with company, year, section, file and ID"],
        ["Index", "Chunk JSONL", "Encode text with BAAI/bge-small-en-v1.5; upsert vectors into Chroma", "Persistent local vector collection"],
        ["Retrieve", "Question and optional filters", "Semantic nearest neighbors plus BM25 keyword rank; reciprocal-rank fusion", "Top passages with citation metadata"],
        ["Tool", "Financial question", "Read-only DuckDB query with validated table name and bound values", "644 reported/derived historical rows, 36 forecasts, 34 outlier candidates"],
        ["Analytics", "Annual SEC facts", "Compare persistence and linear trend with rolling one-step MAE; apply robust median/MAD score", "Forecast and anomaly tables with model/source metadata"],
        ["Answer", "Evidence and tool output", "Configurable Ollama model; citation and number checks", "Answer, sources, tool trace"],
        ["Serve and evaluate", "Question set or user request", "FastAPI, Streamlit, standalone dashboard, health/readiness endpoints, hit@k evaluation and offline checks", "Conversational response, inspectable evidence, dependency status and measurements"],
    ])
    para(doc, "Chroma and BM25 solve different retrieval problems. Embeddings can find conceptually similar wording; BM25 rewards shared terms. Reciprocal-rank fusion combines their ordered results without pretending that their raw scores use the same scale.")

    doc.add_heading("4. Build steps in order", level=1)
    steps = [
        "Create a Python 3.11 virtual environment and install the runtime dependencies declared in pyproject.toml.",
        "Copy .env.example to .env. Add the SEC User-Agent contact email and keep .env private.",
        "Install Ollama. For this laptop, sign in with ollama signin and set OLLAMA_MODEL=gemma4:cloud in .env; this avoids downloading model weights locally. For offline use, select a small local model such as qwen2.5:1.5b and run ollama pull qwen2.5:1.5b; close memory-heavy apps first.",
        "Set the desired ticker list. The default is HPE, MSFT, AAPL, NVDA, and ACN.",
        "Run python -m src.fetch to resolve SEC CIKs, list recent 10-K forms, and download filings with accession-based filenames.",
        "Run python -m src.snapshot to store accession, file size, and SHA-256 for the filing snapshot.",
        "Run python -m src.financial_data to download structured company facts and build historical, forecast, and anomaly tables.",
        "Run python -m src.eda to export company trend plots, per-metric forecast comparisons, anomaly summaries, and data-coverage tables.",
        "Run python -m src.parse to extract section-specific text and save chunks.jsonl.",
        "Inspect a random sample of chunks. Fix parser errors before indexing; bad chunks make retrieval and generated answers unreliable.",
        "Run python -m src.index to create normalized sentence embeddings and persist the Chroma collection.",
        "Run python -m eval.run_eval eval/questions.jsonl after copying and reviewing the seed set. The current metric is retrieval hit@k.",
        "Configure DuckDB path, table names, and schema. Confirm that company, metric, period, and value columns match the adapter.",
        "Start FastAPI with uvicorn api.main:app --reload or run streamlit run app/chat.py. GET /health checks process liveness; GET /ready checks the index, financial database, Ollama service, and configured model. Check citations and tool output on representative prompts.",
        "Run python scripts/serve_showcase.py in another PowerShell window and open the URL in README.md. The page defaults to an honest demo; add the documented API URL query parameter for live conversation, retrieval diagnostics, company comparisons, forecasts, and anomaly review.",
        "Complete held-out evaluation, document failure cases, and only then report measured performance in a portfolio or resume.",
    ]
    for item in steps:
        number(doc, item)

    doc.add_heading("5. Module guide: what each code file does and why", level=1)
    modules = [
        ("pyproject.toml", "Defines the package, Python version, runtime dependencies, and optional development tools. This makes installation reproducible and separates app packages from test/lint packages."),
        ("src/config.py", "Reads settings from environment variables or .env. Credentials and machine-specific paths do not belong in committed code. The sample local default is Qwen2.5 1.5B; this machine's ignored .env selects gemma4:cloud because available local memory is limited. Model choice and request timeout can be changed in .env."),
        ("src/fetch.py", "Calls SEC endpoints with a contact User-Agent, resolves CIKs from the SEC ticker map, finds recent 10-K records, and saves accession-named HTML files. Caching avoids needless downloads."),
        ("src/snapshot.py", "Stores accession, report date, size, and SHA-256 hash for each filing. A hash helps detect whether a source file changed between analyses."),
        ("src/financial_data.py", "Downloads SEC XBRL company facts, standardizes six annual metrics, selects between last-value and linear-trend forecasts with rolling MAE, flags extreme points with median absolute deviation, and writes the tables to DuckDB."),
        ("src/eda.py", "Reads the DuckDB tables and exports financial data coverage, annual revenue and margin trends, forecast MAPE comparisons by metric, and anomaly-candidate plots for review."),
        ("src/parse.py", "Removes scripts and hidden markup, locates Item 1, Item 1A, Item 7, and Item 8, and creates overlapping word chunks. Metadata links every chunk to its company and source file."),
        ("src/index.py", "Encodes each chunk with a local sentence-transformer and writes vectors and metadata to persistent Chroma in small batches."),
        ("src/retrieve.py", "Applies optional company/year filters, searches the vector collection and a BM25 index, then fuses ranks. Fast BM25 word search is the default and does not load the embedding model; deeper hybrid or meaning-based search remains available in the dashboard."),
        ("src/tools.py", "Searches filings and connects to DuckDB read-only, validates configured table identifiers and expected columns, parameterizes company and metric values, and reads historicals, forecasts, or anomalies."),
        ("src/agent.py", "Routes company filing questions to retrieval and financial tools, while keeping greetings, general finance definitions, data-science questions, and project questions conversational. Follow-up turns carry recent chat context; filing questions remain evidence grounded. Fast word search and three evidence passages are the defaults; users can select deeper retrieval in the dashboard."),
        ("src/generate.py", "Prompts Ollama to behave conversationally for general questions and to answer in everyday English for people who are new to finance. It defines key terms, shortens jargon, and explains what company-report passages mean. It cites retrieved IDs, uses tool output as the sole source for financial figures, removes invalid citations, and masks only unsupported figures while preserving the cited explanation."),
        ("api/main.py", "Validates chat history, retrieval mode, and company-comparison requests; exposes /health, /ready, /ask, and read-only /analytics; restricts CORS to the local showcase origin; and returns actionable errors for unavailable Ollama models."),
        ("app/chat.py", "Provides a Streamlit chat page with simple-language guidance, optional company/year filters, expandable source/tool panels, and setup guidance for local model downloads or cloud sign-in."),
        ("showcase.html", "A self-contained HTML/CSS/JavaScript dashboard with themes, animated background, company scope, a plain-English conversational assistant, helpful question examples, beginner-friendly source labels, saved sessions, company comparison, read-only analytics views, filing manifest, and export tools. Search controls are tucked into an optional panel; live API access is opt-in through its URL parameter; demo mode does not invent filing answers."),
        ("scripts/serve_showcase.py", "Serves only showcase.html on localhost. It does not expose the project's .env, DuckDB database, or downloaded filing files."),
        ("eval/run_eval.py", "Measures hit@k and MRR@k for answerable questions, with breakdowns by company and target section. It counts unanswerable prompts but leaves refusal scoring for a separate generated-answer and human-review stage."),
        ("tests/test_core.py", "Checks chunk overlap, section detection, SQL identifier safety, SEC header validation, financial-number guards and redaction behavior, XBRL period selection, derived margins, forecast selection, anomaly-change detection, short greetings, general finance-definition routing, conversational routing, and fast-search model loading without making network calls."),
    ]
    for name, desc in modules:
        doc.add_heading(name, level=2)
        para(doc, desc)

    doc.add_heading("6. Dataset and evaluation design", level=1)
    para(doc, "The plan calls for three annual reports each from four or five companies. The fetcher downloaded fifteen filings: three each for HPE, Microsoft, Apple, NVIDIA, and Accenture. The exact fiscal years differ by issuer, and the accession number is recorded in each filename and in data/filing_manifest.json. SEC company facts add structured annual metrics, source tags, accessions, and values for downstream database tools.")
    table(doc, ["Data source", "Endpoint used"], [
        ["SEC ticker-to-CIK map", "https://www.sec.gov/files/company_tickers.json"],
        ["SEC filing submissions", "https://data.sec.gov/submissions/CIK{CIK10}.json"],
        ["SEC filing archive", "https://www.sec.gov/Archives/edgar/data/{CIK}/{accession_without_dashes}/{primary_document}"],
        ["SEC XBRL company facts", "https://data.sec.gov/api/xbrl/companyfacts/CIK{CIK10}.json"],
    ])
    manifest_path = ROOT / "data/filing_manifest.json"
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        table(doc, ["Ticker", "Report date", "SEC accession"], [
            [item["ticker"], item["report_date"], item["accession"]]
            for item in manifest["files"]
        ])
    table(doc, ["Structured data result", "Snapshot result"], [
        ["Reported and derived annual metrics", "644 rows across five companies"],
        ["One-year forecasts", "36 series; 26 selected last-value and 10 selected linear trend"],
        ["Outlier candidates", "34 flags on unusual year-over-year changes using median absolute deviation"],
    ])
    para(doc, "Each forecast uses at least six annual observations. The pipeline compares a last-value baseline with a linear trend on rolling one-step-ahead holdouts from up to the five most recent historical periods and selects the lower-MAE model. Both candidate MAE and MAPE are stored. The projections and anomaly flags are exploratory analysis, not investment advice or proof of future performance.")
    para(doc, "The seed evaluation file contains 50 prompts: 40 answerable prompts spread across five companies and four 10-K sections, plus 10 unanswerable prompts. These are initial prompts, not a verified gold benchmark. A reviewer must check them against the downloaded filings, write reference answers, and attach exact chunk IDs before treating scores as interview-ready results.")
    table(doc, ["Measurement", "How to interpret it", "Current status"], [
        ["Retrieval hit@5", "A target source section appears in the five retrieved results", "0.825 BM25; 0.925 dense; 0.925 hybrid on seed labels only; not held out"],
        ["Retrieval MRR@5", "How highly the first relevant result is ranked", "Added to the evaluation runner; no verified final benchmark yet"],
        ["Citation validity", "Every citation ID in an answer exists in the evidence returned for that question", "Citation post-check implemented; no run yet"],
        ["Numeric correctness", "Financial figures in an answer are present in DuckDB tool output", "Conservative figure guard implemented; SEC-backed DuckDB builder prepared"],
        ["Unanswerable refusal", "The system declines questions unsupported by the filings", "Prompts prepared; requires model run and review"],
        ["Answer quality", "Human review of completeness, relevance, and source support", "No review results yet"],
        ["Ablation", "Compare keyword, dense, and fused retrieval on the same labels", "Seed run: 0.825 BM25; 0.925 dense and hybrid; labels need verification"],
    ])
    para(doc, "The retrieval values above are preliminary smoke results, not final performance claims. A valid result requires reviewed labels, short reference answers, exact gold chunk IDs, and a held-out split. The seed file also includes unanswerable prompts, but refusal quality has not yet been scored with a human-reviewed answer-quality protocol.")

    doc.add_heading("Exploratory financial analysis", level=2)
    para(doc, "The reproducible analysis module reads the same DuckDB tables as the application and writes coverage tables plus three figures. Revenue is converted to USD billions for readability, while operating margins and forecast errors remain percentages. Missing years reflect reported-fact coverage and are not interpolated.")
    para(doc, "The chart compares median MAPE for two simple annual baselines by metric. The production forecast table selects a model per series by lower walk-forward MAE, so MAPE is shown as a scale-free descriptive comparison rather than the selection rule. Results are based on a short historical window and do not establish future accuracy.")
    analysis_path = ROOT / "reports/analysis_summary.json"
    if analysis_path.exists():
        analysis = json.loads(analysis_path.read_text(encoding="utf-8"))
        table(doc, ["Metric", "Series", "Last-value median MAPE %", "Trend median MAPE %", "Selected series (last / trend)"], [
            [
                row["metric"], str(row["series_count"]),
                f"{row['last_value_median_mape_percent']:.2f}",
                f"{row['linear_trend_median_mape_percent']:.2f}",
                f"{row['selected_last_value']} / {row['selected_linear_trend']}",
            ]
            for row in analysis["forecast_mape_comparison"]
        ])
    para(doc, "The anomaly heatmap summarizes 34 candidate flags from unusual annual changes. Each flag is a prompt for review against the original SEC fact and filing, not a confirmed event or accounting error.")
    for image_path, caption in [
        (ROOT / "reports/figures/financial_trends.png", "Figure 1. Annual reported revenue and operating margin by company."),
        (ROOT / "reports/figures/forecast_model_comparison.png", "Figure 2. Median walk-forward MAPE by metric for two baselines."),
        (ROOT / "reports/figures/anomaly_candidates.png", "Figure 3. Anomaly-candidate counts by company and metric."),
    ]:
        if image_path.exists():
            doc.add_picture(str(image_path), width=Inches(6.8))
            caption_paragraph = doc.add_paragraph(caption)
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption_paragraph.runs:
                run.italic = True
                run.font.size = Pt(9)

    doc.add_heading("7. Data quality, reliability, and safety", level=1)
    for item in [
        "Parser boundary risk: annual reports repeat headings in their tables of contents, and embedded tables can disrupt text order. The parser uses heuristic section heading detection and needs sample review.",
        "Year labeling risk: the initial parser uses SEC report date as year metadata. Confirm fiscal-year labels before building gold questions.",
        "Numeric risk: the DuckDB adapter expects the documented schema. It cannot guess or derive absent values.",
        "Refusal risk: a language model may answer beyond retrieved evidence. Prompt rules and citation filtering help but do not prove factual accuracy.",
        "Data and secrets: raw filings, Chroma data, DuckDB files, and .env are excluded from Git. SEC traffic uses a user-provided contact in the User-Agent.",
        "Decision limit: public company filings are not financial advice. The UI says this directly.",
    ]:
        bullet(doc, item)

    doc.add_heading("8. Run commands and expected artifacts", level=1)
    commands = [
        "py -3.11 -m venv .venv",
        ".\\.venv\\Scripts\\Activate.ps1",
        "pip install -e \".[dev]\"",
        "Copy-Item .env.example .env   # then add SEC contact email",
        "python -m src.fetch",
        "python -m src.snapshot",
        "python -m src.financial_data",
        "python -m src.eda",
        "python -m src.parse",
        "python -m src.index",
        "python -m eval.run_eval eval/questions.jsonl",
        "pytest",
        "uvicorn api.main:app --reload",
        "streamlit run app/chat.py",
    ]
    for command in commands:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(command)
        r.font.name = "Cascadia Mono"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(30, 75, 92)
    table(doc, ["Artifact", "Purpose"], [
        ["data/filings/*.html", "15 downloaded source filings (local, ignored by Git)"],
        ["data/filing_manifest.json", "Accession, file size, and SHA-256 per filing"],
        ["data/annual_financials.jsonl", "Normalized annual SEC facts with source tags and accession numbers"],
        ["data/fincopilot.duckdb", "Read-only finance tool source with 644 fact/derived rows, 36 forecasts, and 34 anomaly candidates"],
        ["data/chunks/chunks.jsonl", "1,519 parsed filing chunks and retrieval metadata"],
        ["data/index/chroma", "Persistent vector index"],
        ["reports/figures and reports/tables", "Financial trend, baseline comparison, anomaly, and coverage outputs"],
        ["eval/questions.seed.jsonl", "50 draft prompts: 40 answerable and 10 unanswerable; needs label review"],
        ["eval/results.seed.json", "Preliminary, unverified section-level retrieval comparison"],
        ["FinCopilot_Data_Science_Project_Report.docx", "This complete build walkthrough and source appendix"],
    ])

    doc.add_heading("9. Portfolio description guidance", level=1)
    para(doc, "A truthful portfolio summary after execution can describe the data source, number of filings, retrieval design, and measured held-out metrics. Replace every placeholder with numbers produced by the run. Do not describe unrun measurements as achievements.")
    p = doc.add_paragraph()
    r = p.add_run("Example wording after validation: “Built a retrieval-augmented financial filing assistant over [N] SEC 10-K filings for five companies; combined BM25 with local sentence embeddings, and measured [X]% hit@5 on [N] human-labeled held-out questions. Added source-ID checks and DuckDB-backed financial tools.”")
    r.italic = True

    doc.add_page_break()
    doc.add_heading("Appendix A Complete source code", level=1)
    para(doc, "The listings below are the full source files present when this report was generated. The project repository is the editable source of truth.")
    source_paths = [
        ROOT / "pyproject.toml", ROOT / ".env.example", ROOT / ".gitignore",
        ROOT / "src/config.py", ROOT / "src/fetch.py", ROOT / "src/snapshot.py", ROOT / "src/financial_data.py", ROOT / "src/eda.py", ROOT / "src/parse.py", ROOT / "src/index.py",
        ROOT / "src/retrieve.py", ROOT / "src/tools.py", ROOT / "src/generate.py", ROOT / "src/agent.py",
        ROOT / "src/__init__.py", ROOT / "api/__init__.py", ROOT / "scripts/build_report.py",
        ROOT / "api/main.py", ROOT / "app/chat.py", ROOT / "showcase.html", ROOT / "scripts/serve_showcase.py", ROOT / "eval/run_eval.py",
        ROOT / "eval/questions.seed.jsonl", ROOT / "eval/results.seed.json", ROOT / "tests/test_core.py",
    ]
    for source_index, path in enumerate(source_paths):
        if path.exists():
            # Let the appendix heading and explanation share the opening page
            # with the first source listing instead of leaving a nearly empty page.
            if source_index > 0 and path.name != "test_core.py":
                doc.add_page_break()
            code_listing(doc, path)

    doc.core_properties.title = "FinCopilot Data Science Project Report"
    doc.core_properties.subject = "Build, design, evaluation plan, and complete source listing"
    doc.core_properties.author = "FinCopilot Project"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
